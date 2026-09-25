"""Ekstrakcja treści z arkuszy XLSX biblioteką openpyxl.

Skoroszyt jest czytany w trybie tylko do odczytu, strumieniowo, więc duży arkusz
nie zajmuje pamięci ponad to, co jest potrzebne na zapis tabeli. Wartości
komórek pochodzą z zapisanych w pliku wyników formuł, a nie z samych formuł:
czytelnik notatnika potrzebuje tego, co widział autor, a nie wzoru. Formuła, której
wynik nie został zapisany w pliku, na przykład w arkuszu wygenerowanym programem
bez silnika obliczeń, dałaby pustą komórkę, więc ich liczba trafia do ostrzeżeń
ekstraktora: utrata treści nie jest cicha.

Każdy arkusz jest nagłówkiem „Arkusz: nazwa” i jedną tabelą, z pierwszym wierszem
jako nagłówkiem kolumn, jak w plikach CSV. Arkusz ukryty jest odczytywany i
oznaczony. Wykresy i obrazy nie są odczytywane, a ich liczba trafia do ostrzeżeń.
Zasady zapisu wartości opisuje `gnb.extractors.arkusze`.

Plik uszkodzony albo zaszyfrowany hasłem, który nie jest archiwum ZIP, kończy się
błędem trwałym z czytelnym komunikatem.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterator
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from gnb.core.model import BlokTresci, DokumentWyekstrahowany
from gnb.core.stale import PoziomPewnosciStruktury, TypZrodla
from gnb.core.wyjatki import BladTrwaly
from gnb.extractors.arkusze import bloki_arkusza, czy_format_walutowy, wartosc_na_tekst
from gnb.extractors.bazowy import PostepEkstrakcji
from gnb.extractors.bloki_markdown import zapisz_bloki_jako_markdown

METODA_EKSTRAKCJI = "xlsx"
FORMATY_XLSX = frozenset({"xlsx", "xlsm"})

KOMUNIKAT_USZKODZONY = (
    "Plik XLSX jest uszkodzony, zaszyfrowany hasłem albo nie jest skoroszytem programu "
    "Excel: nie dało się go odczytać. Skoroszytu zaszyfrowanego hasłem aplikacja nie "
    "otwiera. Zapisz go bez szyfrowania i dodaj ponownie."
)

_BLEDY_ODCZYTU = (
    InvalidFileException,
    zipfile.BadZipFile,
    KeyError,
    ValueError,
    OSError,
    TypeError,
    AttributeError,
)
_FORMAT_DATY = "%Y-%m-%d"


class EkstraktorXlsx:
    """Ekstraktor skoroszytów XLSX: każdy arkusz jako nagłówek i tabela."""

    metoda = METODA_EKSTRAKCJI
    tekst_zawiera_znaczniki = True

    def obsluguje(self, typ_zrodla: TypZrodla, format_zrodla: str) -> bool:
        return typ_zrodla is TypZrodla.PLIK_DOKUMENT and format_zrodla in FORMATY_XLSX

    def wyekstrahuj(
        self,
        identyfikator_zrodla: str,
        bajty: bytes,
        *,
        postep: PostepEkstrakcji | None = None,
    ) -> DokumentWyekstrahowany:
        """Odczytuje wszystkie arkusze skoroszytu. Argument `postep` nie jest używany."""
        try:
            zeszyt = load_workbook(io.BytesIO(bajty), read_only=True, data_only=True)
            zeszyt_formul = load_workbook(io.BytesIO(bajty), read_only=True, data_only=False)
        except _BLEDY_ODCZYTU as blad:
            raise BladTrwaly(KOMUNIKAT_USZKODZONY, identyfikator_zrodla) from blad

        try:
            bloki: list[BlokTresci] = []
            formuly_bez_wartosci = 0
            komorki_walutowe = 0
            for arkusz in zeszyt.worksheets:
                wiersze, formuly, waluty = _wiersze_arkusza(arkusz, zeszyt_formul[arkusz.title])
                formuly_bez_wartosci += formuly
                komorki_walutowe += waluty
                bloki.extend(
                    bloki_arkusza(arkusz.title, wiersze, ukryty=arkusz.sheet_state != "visible")
                )
            wlasciwosci = zeszyt.properties
            tytul = wlasciwosci.title.strip() if wlasciwosci.title else None
            metadane = _metadane(wlasciwosci)
        except _BLEDY_ODCZYTU as blad:
            raise BladTrwaly(KOMUNIKAT_USZKODZONY, identyfikator_zrodla) from blad
        finally:
            zeszyt.close()
            zeszyt_formul.close()

        return DokumentWyekstrahowany(
            identyfikator_zrodla=identyfikator_zrodla,
            tekst=zapisz_bloki_jako_markdown(bloki),
            poziom_pewnosci_struktury=PoziomPewnosciStruktury.WYSOKI,
            metoda_ekstrakcji=METODA_EKSTRAKCJI,
            tytul=tytul or None,
            bloki=bloki,
            metadane=metadane,
            ostrzezenia=_ostrzezenia(bajty, formuly_bez_wartosci, komorki_walutowe),
        )


def _wiersze_arkusza(arkusz: Any, arkusz_formul: Any) -> tuple[list[list[str]], int, int]:
    """Zwraca wiersze arkusza jako tekst, liczbę formuł bez wyniku i liczbę komórek walutowych.

    Wymiary arkusza są zerowane przed odczytem: plik zapisany przez program,
    który zapisuje błędny znacznik wymiarów, na przykład jedną komórkę, dawałby
    w trybie strumieniowym tylko tę komórkę.
    """
    arkusz.reset_dimensions()
    arkusz_formul.reset_dimensions()
    wiersze: list[list[str]] = []
    formuly_bez_wartosci = 0
    komorki_walutowe = 0
    for wiersz, wiersz_formul in zip(
        arkusz.iter_rows(), arkusz_formul.iter_rows(values_only=True), strict=False
    ):
        komorki: list[str] = []
        for komorka, formula in _pary(wiersz, wiersz_formul):
            wartosc = getattr(komorka, "value", None)
            format_liczby = getattr(komorka, "number_format", None)
            komorki.append(wartosc_na_tekst(wartosc, format_liczby))
            if isinstance(wartosc, int | float) and czy_format_walutowy(format_liczby):
                komorki_walutowe += 1
            if wartosc is None and isinstance(formula, str) and formula.startswith("="):
                formuly_bez_wartosci += 1
        wiersze.append(komorki)
    return wiersze, formuly_bez_wartosci, komorki_walutowe


def _pary(wiersz: Any, wiersz_formul: Any) -> Iterator[tuple[Any, Any]]:
    """Paruje komórki wiersza z odpowiadającymi im formułami, dopełniając brak Nonami."""
    formuly = list(wiersz_formul)
    for indeks, komorka in enumerate(wiersz):
        yield komorka, formuly[indeks] if indeks < len(formuly) else None


def _metadane(wlasciwosci: Any) -> dict[str, str]:
    metadane: dict[str, str] = {}
    if wlasciwosci.creator and str(wlasciwosci.creator).strip():
        metadane["autor"] = str(wlasciwosci.creator).strip()
    if wlasciwosci.created is not None:
        metadane["data_publikacji"] = wlasciwosci.created.strftime(_FORMAT_DATY)
    if wlasciwosci.modified is not None:
        metadane["data_aktualizacji"] = wlasciwosci.modified.strftime(_FORMAT_DATY)
    return metadane


def _ostrzezenia(bajty: bytes, formuly_bez_wartosci: int, komorki_walutowe: int) -> list[str]:
    """Zbiera ostrzeżenia o treści, której ekstraktor nie odczytał albo uprościł."""
    ostrzezenia: list[str] = []
    if komorki_walutowe:
        ostrzezenia.append(
            f"W skoroszycie jest {komorki_walutowe} komórek z formatem walutowym. Zapisano "
            "same liczby, bez symbolu waluty, więc przy tych komórkach nie ma jednostki."
        )
    if formuly_bez_wartosci:
        ostrzezenia.append(
            f"W skoroszycie jest {formuly_bez_wartosci} formuł bez zapisanego wyniku, więc ich "
            "komórki są puste. Otwórz plik w programie arkusza, zapisz go ponownie i dodaj "
            "jeszcze raz, żeby wyniki zostały zapisane."
        )
    try:
        with zipfile.ZipFile(io.BytesIO(bajty)) as archiwum:
            nazwy = archiwum.namelist()
    except zipfile.BadZipFile:
        return ostrzezenia
    obrazy = sum(1 for nazwa in nazwy if nazwa.startswith("xl/media/"))
    wykresy = sum(1 for nazwa in nazwy if nazwa.startswith("xl/charts/") and nazwa.endswith(".xml"))
    if obrazy:
        ostrzezenia.append(
            f"Skoroszyt zawiera obrazy ({obrazy}), których treść nie została odczytana."
        )
    if wykresy:
        ostrzezenia.append(
            f"Skoroszyt zawiera wykresy ({wykresy}), których treść nie została odczytana."
        )
    return ostrzezenia
