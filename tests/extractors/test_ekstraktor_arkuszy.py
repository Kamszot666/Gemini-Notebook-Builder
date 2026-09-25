"""Testy ekstraktorów XLSX i XLS oraz wspólnych przekształceń arkuszy."""

from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import pytest
from openpyxl import Workbook

from gnb.core.stale import PoziomPewnosciStruktury, TypZrodla
from gnb.core.wyjatki import BladTrwaly
from gnb.extractors.arkusze import bloki_arkusza, czy_format_walutowy, wartosc_na_tekst
from gnb.extractors.plik_xls import EkstraktorXls
from gnb.extractors.plik_xlsx import EkstraktorXlsx

DANE = Path(__file__).resolve().parents[1] / "dane" / "formaty"


def _skoroszyt(zapisz: object) -> bytes:
    """Buduje XLSX biblioteką openpyxl; przekazana funkcja wypełnia skoroszyt."""
    zeszyt = Workbook()
    zapisz(zeszyt)  # type: ignore[operator]
    bufor = io.BytesIO()
    zeszyt.save(bufor)
    return bufor.getvalue()


# --- wartości komórek ----------------------------------------------------------


@pytest.mark.parametrize(
    ("wartosc", "oczekiwany"),
    [
        (None, ""),
        (True, "prawda"),
        (False, "fałsz"),
        (dt.datetime(2026, 5, 7), "2026-05-07"),
        (dt.datetime(2026, 5, 7, 13, 45, 10), "2026-05-07 13:45:10"),
        (dt.date(2026, 5, 7), "2026-05-07"),
        (dt.time(8, 30), "08:30:00"),
        (7, "7"),
        (3.0, "3"),
        (0.1 + 0.2, "0.3"),
        (1234.5, "1234.5"),
        ("tekst", "tekst"),
    ],
)
def test_wartosc_na_tekst(wartosc: object, oczekiwany: str) -> None:
    assert wartosc_na_tekst(wartosc) == oczekiwany


def test_format_procentowy_zapisuje_procent() -> None:
    assert wartosc_na_tekst(0.25, "0%") == "25%"
    assert wartosc_na_tekst(0.255, "0.0%") == "25.5%"
    assert wartosc_na_tekst(0.25, "0.00") == "0.25"


def test_wykrywanie_formatu_walutowego() -> None:
    assert czy_format_walutowy('#,##0.00\\ "zł"')
    assert czy_format_walutowy("[$€-2] #,##0.00")
    assert czy_format_walutowy("$#,##0")
    assert not czy_format_walutowy("0.00")
    assert not czy_format_walutowy("General")
    assert not czy_format_walutowy(None)


def test_arkusz_bez_tresci_nie_daje_blokow() -> None:
    assert bloki_arkusza("Pusty", [["", ""], []]) == []


def test_arkusz_ukryty_jest_oznaczony_w_naglowku() -> None:
    bloki = bloki_arkusza("Sekret", [["a"], ["b"]], ukryty=True)

    assert bloki[0].tresc == "Arkusz: Sekret (arkusz ukryty)"


# --- XLSX ----------------------------------------------------------------------


def test_xlsx_z_libreoffice_zapisuje_arkusze_daty_i_ostrzega_o_walucie() -> None:
    dokument = EkstraktorXlsx().wyekstrahuj("id", (DANE / "arkusz.xlsx").read_bytes())

    assert dokument.poziom_pewnosci_struktury is PoziomPewnosciStruktury.WYSOKI
    assert "## Arkusz: Sprzedaż" in dokument.tekst
    assert "| Produkt | Cena | Data |" in dokument.tekst
    # Data jest datą, a nie liczbą zależną od stylu komórki.
    assert "2026-05-07" in dokument.tekst
    assert "## Arkusz: Drugi" in dokument.tekst
    assert any("formatem walutowym" in ostrzezenie for ostrzezenie in dokument.ostrzezenia)


def test_xlsx_formula_bez_zapisanego_wyniku_jest_zglaszana_a_nie_ukryta() -> None:
    def wypelnij(zeszyt: Workbook) -> None:
        arkusz = zeszyt.active
        arkusz.append(["a", "b"])
        arkusz.append([1, 2])
        arkusz["C2"] = "=A2+B2"

    dokument = EkstraktorXlsx().wyekstrahuj("id", _skoroszyt(wypelnij))

    assert any(
        "1 formuł bez zapisanego wyniku" in ostrzezenie for ostrzezenie in dokument.ostrzezenia
    )


def test_xlsx_zapisuje_wartosci_logiczne_daty_z_godzina_i_procenty() -> None:
    def wypelnij(zeszyt: Workbook) -> None:
        arkusz = zeszyt.active
        arkusz.append(["Prawda", "Data", "Procent", "Liczba"])
        arkusz.append([True, dt.datetime(2026, 1, 2, 3, 4, 5), 0.25, 0.1 + 0.2])
        arkusz["C2"].number_format = "0%"

    tekst = EkstraktorXlsx().wyekstrahuj("id", _skoroszyt(wypelnij)).tekst

    assert "| prawda | 2026-01-02 03:04:05 | 25% | 0.3 |" in tekst


def test_xlsx_arkusz_ukryty_jest_odczytany_i_oznaczony() -> None:
    def wypelnij(zeszyt: Workbook) -> None:
        zeszyt.active.append(["jawny"])
        zeszyt.active.append(["x"])
        ukryty = zeszyt.create_sheet("Ukryty")
        ukryty.append(["tajny"])
        ukryty.append(["y"])
        ukryty.sheet_state = "hidden"

    tekst = EkstraktorXlsx().wyekstrahuj("id", _skoroszyt(wypelnij)).tekst

    assert "## Arkusz: Ukryty (arkusz ukryty)" in tekst
    assert "tajny" in tekst


def test_xlsx_pomija_puste_arkusze_i_wiersze() -> None:
    def wypelnij(zeszyt: Workbook) -> None:
        zeszyt.active.append(["a"])
        zeszyt.active.append([])
        zeszyt.active.append(["b"])
        zeszyt.create_sheet("Pusty")

    dokument = EkstraktorXlsx().wyekstrahuj("id", _skoroszyt(wypelnij))

    assert "Pusty" not in dokument.tekst
    assert "| a |" in dokument.tekst and "| b |" in dokument.tekst


def test_xlsx_uszkodzony_albo_zaszyfrowany_zglasza_blad_trwaly() -> None:
    # Plik zaszyfrowany hasłem to kontener OLE, a nie archiwum ZIP.
    naglowek_ole = bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 512

    with pytest.raises(BladTrwaly, match="zaszyfrowany hasłem"):
        EkstraktorXlsx().wyekstrahuj("id", naglowek_ole)
    with pytest.raises(BladTrwaly):
        EkstraktorXlsx().wyekstrahuj("id", b"to nie jest skoroszyt")


def test_xlsx_obsluguje_tylko_swoje_formaty() -> None:
    ekstraktor = EkstraktorXlsx()

    assert ekstraktor.obsluguje(TypZrodla.PLIK_DOKUMENT, "xlsx")
    assert ekstraktor.obsluguje(TypZrodla.PLIK_DOKUMENT, "xlsm")
    assert not ekstraktor.obsluguje(TypZrodla.PLIK_DOKUMENT, "xls")


# --- XLS -------------------------------------------------------------------------


def test_xls_z_libreoffice_zapisuje_arkusze_i_daty() -> None:
    dokument = EkstraktorXls().wyekstrahuj("id", (DANE / "arkusz.xls").read_bytes())

    assert "## Arkusz: Sprzedaż" in dokument.tekst
    assert "| Produkt | Cena | Data |" in dokument.tekst
    assert "2026-05-07" in dokument.tekst
    assert "## Arkusz: Drugi" in dokument.tekst
    assert "| a | 3 |" in dokument.tekst
    assert any("formatem walutowym" in ostrzezenie for ostrzezenie in dokument.ostrzezenia)


def test_xls_uszkodzony_zglasza_blad_trwaly() -> None:
    with pytest.raises(BladTrwaly, match="uszkodzony, zaszyfrowany hasłem"):
        EkstraktorXls().wyekstrahuj("id", b"to nie jest skoroszyt")


def test_xls_obsluguje_tylko_xls() -> None:
    ekstraktor = EkstraktorXls()

    assert ekstraktor.obsluguje(TypZrodla.PLIK_DOKUMENT, "xls")
    assert not ekstraktor.obsluguje(TypZrodla.PLIK_DOKUMENT, "xlsx")


def test_xlsx_z_blednym_znacznikiem_wymiarow_jest_czytany_w_calosci() -> None:
    """Program zapisujący plik potrafi podać wymiar jednej komórki, choć arkusz ma ich wiele.

    W trybie strumieniowym taki plik dawałby tylko pierwszą komórkę, więc ekstraktor
    zeruje wymiary przed odczytem.
    """
    import zipfile

    def wypelnij(zeszyt: Workbook) -> None:
        arkusz = zeszyt.active
        arkusz.append(["a", "b", "c"])
        arkusz.append([1, 2, 3])
        arkusz.append([4, 5, 6])

    oryginal = _skoroszyt(wypelnij)
    bufor = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(oryginal)) as zrodlo, zipfile.ZipFile(bufor, "w") as cel:
        for wpis in zrodlo.namelist():
            dane = zrodlo.read(wpis)
            if wpis == "xl/worksheets/sheet1.xml":
                assert b'<dimension ref="A1:C3"/>' in dane
                dane = dane.replace(b'<dimension ref="A1:C3"/>', b'<dimension ref="A1"/>')
            cel.writestr(wpis, dane)

    tekst = EkstraktorXlsx().wyekstrahuj("id", bufor.getvalue()).tekst

    assert "| a | b | c |" in tekst
    assert "| 4 | 5 | 6 |" in tekst
