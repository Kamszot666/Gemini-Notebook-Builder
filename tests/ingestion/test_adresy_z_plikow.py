"""Testy automatycznego wyłuskiwania adresów z plików TXT, MD i DOCX.

Decyzja użytkownika: adresy są pobierane automatycznie wyłącznie z plików TXT,
MD i DOCX, i tylko adresy jawne, czyli zapisane w widocznym tekście. Odnośnik
ukryty pod innym tekstem oraz pliki HTML i MHTML nie dają żadnego nowego źródła.
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from gnb.core.konfiguracja import Konfiguracja
from gnb.ingestion.adresy_z_plikow import dolacz_adresy_znalezione, przyjmij_plik_z_adresami
from gnb.ingestion.lista_url import (
    adresy_z_pliku_z_limitem,
    adresy_znalezione_w_pliku,
    rozpoznaj_liste_adresow_w_pliku,
)

_MOMENT = datetime(2026, 9, 26, 9, 0, tzinfo=UTC)
_JAWNY = "https://jawny.example/artykul"
_UKRYTY = "https://ukryty.example/pod-slowem"


def _docx_z_adresem_jawnym_i_ukrytym(sciezka: Path) -> Path:
    dokument = docx.Document()
    dokument.add_paragraph(f"Zobacz {_JAWNY} oraz")
    akapit = dokument.add_paragraph()
    identyfikator = dokument.part.relate_to(_UKRYTY, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    odnosnik = OxmlElement("w:hyperlink")
    odnosnik.set(qn("r:id"), identyfikator)
    przebieg = OxmlElement("w:r")
    tekst = OxmlElement("w:t")
    tekst.text = "strona gminy"
    przebieg.append(tekst)
    odnosnik.append(przebieg)
    akapit._p.append(odnosnik)
    dokument.save(str(sciezka))
    return sciezka


def _docx_z_akapitami(sciezka: Path, akapity: list[str]) -> Path:
    dokument = docx.Document()
    for akapit in akapity:
        dokument.add_paragraph(akapit)
    dokument.save(str(sciezka))
    return sciezka


def test_docx_daje_adres_jawny_a_ukrytego_pod_slowem_pomija(tmp_path: Path) -> None:
    plik = _docx_z_adresem_jawnym_i_ukrytym(tmp_path / "dokument.docx")

    adresy = adresy_znalezione_w_pliku(plik)

    assert [adres.podany for adres in adresy] == [_JAWNY]


def test_docx_bez_adresow_jawnych_nie_daje_zadnego(tmp_path: Path) -> None:
    plik = _docx_z_akapitami(tmp_path / "dokument.docx", ["Zwykły tekst bez odnośników."])

    assert adresy_znalezione_w_pliku(plik) == ()


def test_docx_zlozony_z_samych_adresow_jest_lista_zrodel(tmp_path: Path) -> None:
    plik = _docx_z_akapitami(
        tmp_path / "lista.docx", ["https://przyklad.pl/a", "https://przyklad.pl/b"]
    )

    podsumowanie = rozpoznaj_liste_adresow_w_pliku(plik)

    assert podsumowanie is not None
    assert [adres.podany for adres in podsumowanie.adresy] == [
        "https://przyklad.pl/a",
        "https://przyklad.pl/b",
    ]


def test_docx_z_tekstem_i_adresem_nie_jest_lista(tmp_path: Path) -> None:
    plik = _docx_z_adresem_jawnym_i_ukrytym(tmp_path / "dokument.docx")

    assert rozpoznaj_liste_adresow_w_pliku(plik) is None


def test_uszkodzony_docx_nie_zatrzymuje_niczego(tmp_path: Path) -> None:
    plik = tmp_path / "uszkodzony.docx"
    plik.write_bytes(b"to nie jest dokument docx")

    assert adresy_znalezione_w_pliku(plik) == ()
    assert rozpoznaj_liste_adresow_w_pliku(plik) is None


def test_md_pomija_odnosnik_ukryty_pod_tekstem_i_bierze_adres_jawny(tmp_path: Path) -> None:
    plik = tmp_path / "notatka.md"
    plik.write_text(
        "# Źródła\n\n"
        f"Zobacz [strona gminy]({_UKRYTY}) oraz {_JAWNY}.\n"
        "Jeszcze [https://tekst.example/a](https://tekst.example/a) i [inny][odn].\n\n"
        "[odn]: https://definicja.example/b\n",
        encoding="utf-8",
    )

    adresy = adresy_znalezione_w_pliku(plik)

    assert [adres.podany for adres in adresy] == [_JAWNY, "https://tekst.example/a"]


def test_html_i_mhtml_nie_daja_zadnych_adresow_ani_listy(tmp_path: Path) -> None:
    tresc_html = (
        "<html><body><p>Zobacz https://tekst.example/a</p>"
        '<a href="https://odnosnik.example/b">https://odnosnik.example/b</a></body></html>'
    )
    mhtml = (
        "MIME-Version: 1.0\r\n"
        'Content-Type: multipart/related; boundary="granica"\r\n\r\n'
        "--granica\r\n"
        "Content-Type: text/html; charset=utf-8\r\n\r\n"
        f"{tresc_html}\r\n--granica--\r\n"
    )
    for nazwa, tresc in (
        ("strona.html", tresc_html),
        ("strona.htm", tresc_html),
        ("strona.xhtml", tresc_html),
        ("strona.mhtml", mhtml),
        ("strona.mht", mhtml),
    ):
        plik = tmp_path / nazwa
        plik.write_text(tresc, encoding="utf-8")

        assert adresy_znalezione_w_pliku(plik) == (), nazwa
        assert rozpoznaj_liste_adresow_w_pliku(plik) is None, nazwa
        przyjecie = przyjmij_plik_z_adresami(plik, _MOMENT, Konfiguracja())
        assert len(przyjecie.pozycje) == 1, nazwa
        assert przyjecie.adresy_znalezione == (), nazwa


def test_dokladnie_limit_adresow_z_docx_i_md_przechodzi_a_ponad_limit_nie(
    tmp_path: Path,
) -> None:
    def adresy(liczba: int) -> list[str]:
        return [f"https://przyklad.pl/strona{numer}" for numer in range(liczba)]

    pliki = {
        "docx": lambda liczba: _docx_z_akapitami(
            tmp_path / f"a{liczba}.docx", ["Zobacz: " + " oraz ".join(adresy(liczba))]
        ),
        "md": lambda liczba: _md(
            tmp_path / f"a{liczba}.md", "Zobacz: " + " oraz ".join(adresy(liczba))
        ),
    }
    for rodzaj, zrob in pliki.items():
        w_limicie = adresy_z_pliku_z_limitem(zrob(200), 200)
        ponad = adresy_z_pliku_z_limitem(zrob(201), 200)

        assert w_limicie.przekroczono_limit is False, rodzaj
        assert len(w_limicie.adresy) == 200, rodzaj
        assert ponad.przekroczono_limit is True, rodzaj
        assert ponad.adresy == (), rodzaj
        assert ponad.liczba_znalezionych == 201, rodzaj


def _md(sciezka: Path, tresc: str) -> Path:
    sciezka.write_text(tresc, encoding="utf-8")
    return sciezka


def test_przyjecie_docx_daje_plik_jako_zrodlo_i_jeden_adres_do_dolaczenia(
    tmp_path: Path,
) -> None:
    plik = _docx_z_adresem_jawnym_i_ukrytym(tmp_path / "dokument.docx")
    konfiguracja = Konfiguracja()

    przyjecie = przyjmij_plik_z_adresami(plik, _MOMENT, konfiguracja)
    pozycje = list(przyjecie.pozycje)
    dolacz_adresy_znalezione(pozycje, list(przyjecie.adresy_znalezione), _MOMENT, konfiguracja)

    assert przyjecie.jest_lista_adresow is False
    assert len(pozycje) == 2
    assert pozycje[0].format_zrodla == "docx"
    assert pozycje[1].wskazane_jawnie is False
    assert pozycje[1].adres_kanoniczny is not None
    assert "jawny.example" in pozycje[1].adres_kanoniczny


def test_przyjecie_pliku_ponad_limit_nie_dodaje_adresow_i_niesie_ostrzezenie(
    tmp_path: Path,
) -> None:
    plik = _md(
        tmp_path / "wiele.md", "Zobacz " + " ".join(f"https://przyklad.pl/s{n}" for n in range(4))
    )
    konfiguracja = Konfiguracja(limit_adresow_z_pliku=3)

    przyjecie = przyjmij_plik_z_adresami(plik, _MOMENT, konfiguracja)

    assert przyjecie.adresy_znalezione == ()
    assert przyjecie.ostrzezenie is not None
    assert "limit_adresow_z_pliku" in przyjecie.ostrzezenie
    assert len(przyjecie.pozycje) == 1
